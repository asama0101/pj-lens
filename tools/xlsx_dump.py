# -*- coding: utf-8 -*-
"""使い方: python xlsx_dump.py <file.xlsx>

xlsx の全シートを JSON（UTF-8）で標準出力へ出す。
  {"sheets": [{"name": str, "rows": [[セル, ...], ...], "indents": [int, ...], "truncated": bool, "hidden": bool}]}
- 結合セルは、左上の値を結合範囲のすべてのセルへ展開する（上位項目が縦に結合された WBS を読むため）
- 日付は YYYY-MM-DD（時刻があれば YYYY-MM-DDTHH:MM）。整数に見える数値は整数にする
- 文字列は前後の空白（全角を含む）を除き、空は null
- indents は行ごとのインデント量（各行の最初の値のあるセルの、先頭の空白の数と Excel のインデント設定の大きい方）。
  インデントで階層を表す WBS を読むために出す（文字列からは先頭の空白を除くため）
- 1 シート最大 2000 行（超えたら truncated: true）
"""
import datetime
import json
import sys

from openpyxl import load_workbook

MAX_ROWS = 2000


def cell_value(v):
    if isinstance(v, datetime.datetime):
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%dT%H:%M")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def leading_ws(v):
    return len(v) - len(v.lstrip()) if isinstance(v, str) else 0


def sheet_rows(ws):
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        top = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row <= len(grid) and rng.min_col <= len(grid[rng.min_row - 1]) else None
        for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
            for c in range(rng.min_col - 1, min(rng.max_col, len(grid[r]))):
                grid[r][c] = top
    rows = [[cell_value(v) for v in row] for row in grid]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    indents = []
    for r, row in enumerate(rows):
        c = next((i for i, v in enumerate(row) if v is not None), None)
        if c is None:
            indents.append(0)
            continue
        excel = int(getattr(ws.cell(row=r + 1, column=c + 1).alignment, "indent", 0) or 0)
        indents.append(max(leading_ws(grid[r][c]), excel))
    truncated = len(rows) > MAX_ROWS
    return rows[:MAX_ROWS], indents[:MAX_ROWS], truncated


def main(path):
    sys.stdout.reconfigure(encoding="utf-8")
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows, indents, truncated = sheet_rows(ws)
        sheets.append({"name": ws.title, "rows": rows, "indents": indents, "truncated": truncated, "hidden": ws.sheet_state != "visible"})
    json.dump({"sheets": sheets}, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.stderr.write("使い方: python xlsx_dump.py <file.xlsx>\n")
        sys.exit(2)
    try:
        main(sys.argv[1])
    except Exception as e:  # 壊れたファイルなど。理由を標準エラーへ出して終了コード 1
        sys.stderr.write(f"読み取れません: {e}\n")
        sys.exit(1)
